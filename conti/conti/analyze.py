"""Run a question PLAN: a short list of steps over Conti's data, each giving a table, a chart or a network,
always with the rows behind it and how much of it rests on checked data.

plan = {
  "title": "…", "question": "…",
  "manuscripts": ["MPAP_4"] | [],                      # default: all
  "steps": [ step, … ]
}
Steps (each has a unique "id"; "from" names an earlier step):

  {"id", "op": "transactions", "kind": "debt" | "payment" | "all",
         "where": {occupation, debtor, creditor (text contained), has_occupation (bool), debtor_type "person"|"institution",
                   year [a, b], currency "florins"|"lire" (has a sum in it), has_sum (bool),
                   checked_sums_only, approved_tags_only, tagged_creditor_only, exclude_problem_sums (bool)}}
        → one row per debt / payment (see ledger.py)

  {"id", "op": "aggregate", "from": <transactions step>, "by": "occupation"|"debtor"|"creditor"|"year"|"page"|"manuscript"|"debtor_type",
         "currency": "florins"|"lire"|"both" (default both), "rank_by": "median"|"total"|"mean"|"max"|"count" (default median),
         "split": "equal"|"full" (joint debts: share the sum among debtors, or count it in full for each; default equal),
         "over": null | "year" | "page" | "manuscript"   (consistency: rank the groups again in each period),
         "min_n": 1, "top": 25, "chart": {"type": "bar", "metric": "median"} (optional)}
        → one row per group: transactions, people, and per currency n / total / mean / median / max, plus
          checked: how many of its sums Laura checked; with "over": a period × group matrix and, per group,
          periods present, periods ranked first, mean rank.

  {"id", "op": "network", "from": <transactions step>, "top": 50, "rank_by": "count"|"florins"|"lire",
         "include_counterparties": true}
        → directed graph debtor → creditor (who owed whom), edges weighted by number of debts and sums owed per currency;
          payments (kind "payment" rows) are counted apart as paid_n / paid_florins / paid_lire, never added to what is owed;
          nodes ranked by involvement; top N kept plus everyone they owed or who owed them.

  {"id", "op": "query", "query": {…Conti query…}, "group_by": "person" | … }  → rows (and counts) from query.py

run(plan) → {"plan", "steps": {id: result}, "coverage": {...}, "caveats": [...], "t"}
"""
from __future__ import annotations

import csv
import io
import statistics
import zipfile

from . import data, extract, ledger, query, source

OPS = ("transactions", "aggregate", "network", "query")
BY = ("occupation", "debtor", "creditor", "year", "page", "manuscript", "debtor_type", "title")
CUR = ("florins", "lire")


class PlanError(ValueError):
    pass


def _fold(s):
    return query._fold(s)


# ------------------------------------------------------------------ transactions
def _tx_ok(t, w):
    if not w:
        return True
    if w.get("occupation") and not any(_fold(w["occupation"]) in _fold(o) for d in t["debtors"] for o in d["occupations"]):
        return False
    if w.get("debtor") and not any(_fold(w["debtor"]) in _fold(d["name"]) for d in t["debtors"]):
        return False
    if w.get("creditor") and _fold(w["creditor"]) not in _fold(t["creditor"]["name"]):
        return False
    if w.get("has_occupation") is not None and bool(any(d["occupations"] for d in t["debtors"])) != bool(w["has_occupation"]):
        return False
    if w.get("debtor_type") and not any(d["type"] == w["debtor_type"] for d in t["debtors"]):
        return False
    if w.get("has_debtor") is not None and bool(t["debtors"]) != bool(w["has_debtor"]):
        return False
    if w.get("year"):
        a, b = (list(w["year"]) + [None, None])[:2]
        if t["year"] is None or (a and t["year"] < int(a)) or (b and t["year"] > int(b)):
            return False
    if w.get("currency") in CUR and t[w["currency"]] is None:
        return False
    if w.get("has_sum") is not None and (t["florins"] is not None or t["lire"] is not None) != bool(w["has_sum"]):
        return False
    if w.get("checked_sums_only") and not t["sum_checked"]:
        return False
    if w.get("approved_tags_only") and not t["tags_checked"]:
        return False
    if w.get("tagged_creditor_only") and t["creditor"]["basis"] != "tagged":
        return False
    if w.get("exclude_problem_sums") and t["sum_problems"]:
        return False
    return True


def _transactions(step, plan, ctx):
    key = tuple(plan.get("manuscripts") or [])
    if key not in ctx["tx"]:
        ctx["tx"][key] = ledger.transactions(list(key) or None)
    kind = step.get("kind", "debt")
    rows = [t for t in ctx["tx"][key] if (kind == "all" or t["kind"] == kind) and _tx_ok(t, step.get("where"))]
    return {"type": "transactions", "rows": rows, "count": len(rows),
            "counts": {"with_debtor": sum(1 for t in rows if t["debtors"]),
                       "with_sum": sum(1 for t in rows if t["florins"] is not None or t["lire"] is not None),
                       "sum_checked": sum(1 for t in rows if t["sum_checked"]),
                       "sum_problems": sum(1 for t in rows if t["sum_problems"]),
                       "tags_checked": sum(1 for t in rows if t["tags_checked"]),
                       "creditor_tagged": sum(1 for t in rows if t["creditor"]["basis"] == "tagged"),
                       "creditor_assumed": sum(1 for t in rows if t["creditor"]["basis"].startswith("assumed")),
                       "with_occupation": sum(1 for t in rows if any(d["occupations"] for d in t["debtors"])),
                       "dated": sum(1 for t in rows if t["year"])}}


# ------------------------------------------------------------------ aggregate
def _keys(t, by):
    """[(group value, share)] for one transaction."""
    ds = t["debtors"]
    if by == "occupation":
        occs = sorted({o for d in ds for o in d["occupations"]}) or ["(no occupation recorded)"]
        return [(o, 1.0) for o in occs], ds
    if by == "debtor":
        return [(d["name"], 1.0) for d in ds] or [("(no debtor tagged)", 1.0)], ds
    if by == "debtor_type":
        return [(x, 1.0) for x in sorted({d["type"] for d in ds})] or [("(no debtor tagged)", 1.0)], ds
    if by == "creditor":
        return [(t["creditor"]["name"] + ("" if t["creditor"]["basis"] == "tagged" else " (assumed)"), 1.0)], ds
    if by == "year":
        return [(t["year"] or "(undated)", 1.0)], ds
    if by in ("page", "manuscript"):
        return [(t[by], 1.0)], ds
    raise PlanError(f"aggregate: unknown 'by' {by!r} (use one of {', '.join(BY)})")


def _share(t, by, value, split):
    if split == "full":
        return 1.0
    ds = t["debtors"]
    if by == "debtor" and len(ds) > 1:
        return 1 / len(ds)
    if by == "occupation" and len(ds) > 1:
        holders = [d for d in ds if value in d["occupations"]]
        return (len(holders) or 1) / len(ds)
    return 1.0


def _group(rows, by, split, curs):
    g = {}
    for t in rows:
        keys, ds = _keys(t, by)
        for v, _ in keys:
            b = g.setdefault(v, {"value": v, "tx": 0, "people": set(), "ids": [], "checked": 0, "problems": 0,
                                 **{c: [] for c in curs}})
            b["tx"] += 1; b["ids"].append(t["id"])
            b["checked"] += t["sum_checked"]; b["problems"] += t["sum_problems"]
            for d in ds:
                if by == "debtor" and d["name"] != v:
                    continue
                if by != "occupation" or v in d["occupations"] or v == "(no occupation recorded)":
                    b["people"].add(d["name"])
            sh = _share(t, by, v, split)
            for c in curs:
                if t[c] is not None:
                    b[c].append(t[c] * sh)
    out = []
    for b in g.values():
        r = {"value": b["value"], "transactions": b["tx"], "people": len(b["people"]), "people_names": sorted(b["people"]),
             "sums_checked": b["checked"], "sums_flagged": b["problems"], "ids": b["ids"]}
        for c in curs:
            s = ledger.stats(b[c])
            for k, x in s.items():
                r[f"{c}_{k}"] = x
        out.append(r)
    return out


def _rank_key(rank_by, curs):
    c = curs[0]
    if rank_by == "count":
        return lambda r: (r["transactions"], r.get(f"{c}_total") or 0)
    return lambda r: (r.get(f"{c}_{rank_by}") if r.get(f"{c}_{rank_by}") is not None else -1, r["transactions"])


def _aggregate(step, plan, ctx):
    src = ctx["steps"].get(step.get("from"))
    if not src or src["type"] != "transactions":
        raise PlanError("aggregate needs 'from': the id of an earlier transactions step")
    by = step.get("by", "occupation")
    cur = step.get("currency", "both")
    curs = list(CUR) if cur == "both" else [cur]
    if any(c not in CUR for c in curs):
        raise PlanError("currency must be florins, lire or both")
    rank_by = step.get("rank_by", "median")
    split = step.get("split", "equal")
    min_n = int(step.get("min_n") or 1)
    rows = _group(src["rows"], by, split, curs)
    # min_n counts sums in the ranking currency (a median of one sum is not a pattern)
    c0 = curs[0]
    rows = [r for r in rows if (r["transactions"] if rank_by == "count" else r.get(f"{c0}_n", 0)) >= min_n]
    rows.sort(key=_rank_key(rank_by, curs), reverse=True)
    res = {"type": "table", "by": by, "currencies": curs, "rank_by": rank_by, "split": split, "min_n": min_n,
           "rows": rows[: int(step.get("top") or 25)], "groups_total": len(rows), "chart": step.get("chart")}
    over = step.get("over")
    if over:
        periods = {}
        for t in src["rows"]:
            p = t["year"] if over == "year" else t[over]
            if p is None:
                p = "(undated)"
            periods.setdefault(p, []).append(t)
        keep = {r["value"] for r in res["rows"]}
        matrix, ranks = [], {}
        for p in sorted(periods, key=str):
            gr = [r for r in _group(periods[p], by, split, curs)          # "(no occupation recorded)" etc. are not ranked
                  if not str(r["value"]).startswith("(")
                  and (r["transactions"] if rank_by == "count" else r.get(f"{c0}_n", 0)) >= 1]
            gr.sort(key=_rank_key(rank_by, curs), reverse=True)
            for i, r in enumerate(gr, 1):
                ranks.setdefault(r["value"], []).append((p, i, len(gr)))
                if r["value"] in keep:
                    matrix.append({"period": p, "value": r["value"], "rank": i, "of": len(gr),
                                   "metric": r["transactions"] if rank_by == "count" else r.get(f"{c0}_{rank_by}"),
                                   "transactions": r["transactions"]})
        cons = []
        for v in keep:
            if str(v).startswith("("):
                continue
            rk = ranks.get(v, [])
            cons.append({"value": v, "periods_present": len(rk), "periods_total": len(periods),
                         "periods_first": sum(1 for _, i, _ in rk if i == 1),
                         "periods_top3": sum(1 for _, i, _ in rk if i <= 3),
                         "mean_rank": round(statistics.mean(i for _, i, _ in rk), 2) if rk else None})
        cons.sort(key=lambda x: (-x["periods_first"], x["mean_rank"] or 99))
        res["over"] = over
        res["periods"] = [str(p) for p in sorted(periods, key=str)]
        res["matrix"] = matrix
        res["consistency"] = cons
    return res


# ------------------------------------------------------------------ network
def _network(step, plan, ctx):
    src = ctx["steps"].get(step.get("from"))
    if not src or src["type"] != "transactions":
        raise PlanError("network needs 'from': the id of an earlier transactions step")
    nodes, edges = {}, {}

    def node(name, typ, basis=None):
        n = nodes.setdefault(name, {"id": name, "label": name, "type": typ, "owed_n": 0, "owed_florins": 0.0, "owed_lire": 0.0,
                                    "owed_to_n": 0, "owed_to_florins": 0.0, "owed_to_lire": 0.0, "occupations": set(),
                                    "assumed": False})
        if basis and basis.startswith("assumed"):
            n["assumed"] = True
        return n

    for t in src["rows"]:
        if not t["debtors"]:
            continue
        cr = t["creditor"]
        c = node(cr["name"], "creditor", cr["basis"])
        k = len(t["debtors"])
        for d in t["debtors"]:
            n = node(d["name"], d["type"])
            n["occupations"].update(d["occupations"])
            f = (t["florins"] or 0) / k; li = (t["lire"] or 0) / k
            e = edges.setdefault((d["name"], cr["name"]), {"source": d["name"], "target": cr["name"], "n": 0, "florins": 0.0,
                                                          "lire": 0.0, "paid_n": 0, "paid_florins": 0.0, "paid_lire": 0.0,
                                                          "ids": [], "basis": cr["basis"], "kinds": set()})
            e["ids"].append(t["id"]); e["kinds"].add(t["kind"])
            if t["kind"] == "payment":          # repayments are not new debts: counted apart, never added to what is owed
                n["paid_n"] = n.get("paid_n", 0) + 1; n["paid_florins"] = n.get("paid_florins", 0) + f; n["paid_lire"] = n.get("paid_lire", 0) + li
                e["paid_n"] += 1; e["paid_florins"] += f; e["paid_lire"] += li
                continue
            n["owed_n"] += 1; n["owed_florins"] += f; n["owed_lire"] += li
            c["owed_to_n"] += 1; c["owed_to_florins"] += f; c["owed_to_lire"] += li
            e["n"] += 1; e["florins"] += f; e["lire"] += li
    rb = step.get("rank_by", "count")

    def score(n):
        if rb == "florins":
            return n["owed_florins"] + n["owed_to_florins"]
        if rb == "lire":
            return n["owed_lire"] + n["owed_to_lire"]
        return n["owed_n"] + n["owed_to_n"] + n.get("paid_n", 0)
    ranked = sorted(nodes.values(), key=lambda n: (score(n), n["owed_florins"] + n["owed_to_florins"]), reverse=True)
    top = int(step.get("top") or 50)
    keep = {n["id"] for n in ranked[:top]}
    for i, n in enumerate(ranked, 1):
        n["rank"] = i; n["top"] = n["id"] in keep
    if step.get("include_counterparties", True):
        keep |= {e["target"] for e in edges.values() if e["source"] in keep} | {e["source"] for e in edges.values() if e["target"] in keep}
    ns = []
    for n in ranked:
        if n["id"] in keep:
            n = dict(n); n["occupations"] = sorted(n["occupations"])
            for k in ("paid_n", "paid_florins", "paid_lire"):
                n.setdefault(k, 0)
            for k in ("owed_florins", "owed_lire", "owed_to_florins", "owed_to_lire", "paid_florins", "paid_lire"):
                n[k] = round(n[k], 2)
            n["role"] = "borrower & lender" if (n["owed_n"] or n["paid_n"]) and n["owed_to_n"] else ("lender" if n["owed_to_n"] else "borrower")
            ns.append(n)
    es = []
    for e in edges.values():
        if e["source"] in keep and e["target"] in keep:
            e = dict(e); e["kinds"] = sorted(e["kinds"])
            for k in ("florins", "lire", "paid_florins", "paid_lire"):
                e[k] = round(e[k], 2)
            es.append(e)
    lenders = [n for n in ns if n["owed_to_n"]]
    return {"type": "network", "nodes": ns, "edges": es, "rank_by": rb, "top": top,
            "summary": {"people_total": len(nodes), "shown": len(ns), "edges": len(es), "lenders": len(lenders),
                        "edges_assumed": sum(1 for e in es if e["basis"].startswith("assumed")),
                        "edges_tagged": sum(1 for e in es if e["basis"] == "tagged")}}


# ------------------------------------------------------------------ query pass-through
def _query(step, plan, ctx):
    q = dict(step.get("query") or {})
    if plan.get("manuscripts") and not q.get("manuscripts"):
        q["manuscripts"] = plan["manuscripts"]
    res = query.run(q)
    out = {"type": "query", "query": q, "describe": query.describe(q), "unit": res["unit"], "count": res["count"],
           "rows": res["rows"]}
    if step.get("group_by"):
        out["groups"] = query.group(res["rows"], step["group_by"], res["unit"])
        out["group_by"] = step["group_by"]
    return out


# ------------------------------------------------------------------ coverage and caveats
def coverage(mss) -> dict:
    out = {"manuscripts": []}
    for cid in mss:
        st = data.settings(cid)
        pm = source.pages(cid)
        recs = extract.records(cid, st["statuses"])
        tagged_pages = {r["page"] for r in recs if r["tags"]}
        s = extract.summary(recs)
        out["manuscripts"].append({"id": cid, "pages_total": len(pm), "pages_included": s["pages"],
                                   "statuses": st["statuses"], "pages_tagged": len(tagged_pages),
                                   "tags": s["tags"], "tags_approved": s["tags_approved"],
                                   "sums": s["sums"], "sums_checked": s["sums_confirmed"], "entries": s["entries"],
                                   "account_holder": st.get("account_holder") or "", "date_style": st["date_style"]})
    return out


def _caveats(res) -> list[str]:
    cv = []
    txs = [r for r in res["steps"].values() if r["type"] == "transactions"]
    if txs:
        c = txs[0]["counts"]; n = txs[0]["count"]
        cv.append("“De dare” means the person owed the estate; the book rarely says whether the debt is a loan, rent or a purchase.")
        if c["sum_checked"] < max(1, n) * 0.5:
            cv.append(f"Sums are mostly machine-read: {c['sum_checked']} of {c['with_sum']} sums used were checked by you; "
                      f"{c['sum_problems']} carry a rule warning.")
        if c["creditor_assumed"]:
            cv.append(f"{c['creditor_assumed']} of {n} rows have no creditor named; the account holder is assumed (Settings).")
        if c["tags_checked"] < n:
            cv.append(f"People and occupations come from tags: {c['tags_checked']} of {n} rows rest only on tags you approved.")
        cv.append("Florins and lire are kept apart (no exchange rate); ƒ ß 𝒹 a fiorino counted at 29 soldi to the florin.")
    for m in res["coverage"]["manuscripts"]:
        if m["pages_tagged"] < m["pages_included"]:
            cv.append(f"{m['id']}: only {m['pages_tagged']} of {m['pages_included']} included pages are tagged "
                      f"({m['pages_total']} pages in the manuscript), so people and occupations are partial.")
    return cv


# ------------------------------------------------------------------ run
def validate(plan: dict) -> None:
    if not isinstance(plan, dict) or not isinstance(plan.get("steps"), list) or not plan["steps"]:
        raise PlanError("a plan needs a list of steps")
    ids = set()
    for s in plan["steps"]:
        if s.get("op") not in OPS:
            raise PlanError(f"unknown op {s.get('op')!r} (use {', '.join(OPS)})")
        if not s.get("id") or s["id"] in ids:
            raise PlanError("each step needs a unique id")
        if s.get("from") and s["from"] not in ids:
            raise PlanError(f"step {s['id']}: 'from' {s['from']!r} is not an earlier step")
        ids.add(s["id"])
    known = {c["id"] for c in source.collections()}
    bad = [m for m in plan.get("manuscripts") or [] if m not in known]
    if bad:
        raise PlanError(f"unknown manuscripts: {', '.join(bad)}")


def run(plan: dict) -> dict:
    validate(plan)
    ctx = {"tx": {}, "steps": {}}
    fn = {"transactions": _transactions, "aggregate": _aggregate, "network": _network, "query": _query}
    for s in plan["steps"]:
        ctx["steps"][s["id"]] = {**fn[s["op"]](s, plan, ctx), "id": s["id"], "op": s["op"], "title": s.get("title", "")}
    mss = plan.get("manuscripts") or [c["id"] for c in source.collections()]
    res = {"plan": plan, "steps": ctx["steps"], "coverage": coverage(mss), "t": data.now()}
    res["caveats"] = _caveats(res)
    return res


# ------------------------------------------------------------------ a compact view for Claude to write the answer from
def digest(res: dict, rows=30) -> dict:
    out = {"coverage": res["coverage"], "caveats": res["caveats"], "steps": {}}
    for k, s in res["steps"].items():
        d = {kk: v for kk, v in s.items() if kk not in ("rows", "nodes", "edges", "matrix")}
        if s["type"] == "transactions":
            d["sample"] = [{x: t[x] for x in ("id", "page", "lines", "florins", "lire", "display", "date", "sum_checked", "sum_problems")}
                           | {"debtors": [(dd["name"], dd["occupations"]) for dd in t["debtors"]], "creditor": t["creditor"]["name"]}
                           for t in s["rows"][:rows]]
        elif s["type"] == "table":
            d["rows"] = [{kk: v for kk, v in r.items() if kk != "ids"} | {"evidence_ids": r["ids"][:12]} for r in s["rows"][:rows]]
            if s.get("matrix"):
                d["matrix"] = s["matrix"][:200]
        elif s["type"] == "network":
            d["nodes"] = s["nodes"][:60]
            d["edges"] = [{kk: v for kk, v in e.items() if kk != "ids"} for e in sorted(s["edges"], key=lambda e: -e["n"])[:80]]
        elif s["type"] == "query":
            d["sample"] = [{x: r.get(x) for x in ("key", "page", "line", "text", "sum")} for r in s["rows"][:rows]]
            if s.get("groups"):
                d["groups"] = s["groups"][:rows]
        out["steps"][k] = d
    return out


# ------------------------------------------------------------------ exports
def _tx_flat(t):
    return {"Id": t["id"], "Kind": t["kind"], "Manuscript": t["manuscript"], "Page": t["page"],
            "Lines": ", ".join(map(str, t["lines"])), "Debtors": "; ".join(d["name"] for d in t["debtors"]),
            "Occupations": "; ".join(sorted({o for d in t["debtors"] for o in d["occupations"]})),
            "Occupation basis": "; ".join(sorted({d["occupation_basis"] for d in t["debtors"] if d["occupation_basis"]})),
            "Creditor": t["creditor"]["name"], "Creditor basis": t["creditor"]["basis"],
            "Florins": t["florins"], "Lire": t["lire"], "Sum (read)": t["display"], "Sum (as written)": t["sum_written"],
            "Sum checked": "yes" if t["sum_checked"] else "no", "Sum warning": "yes" if t["sum_problems"] else "",
            "Tags checked": "yes" if t["tags_checked"] else "no", "Date (modern)": t["date"],
            "Year (as written)": t["year_written"] or "", "Year inferred": "yes" if t["year_inferred"] else "",
            "Text": t["text"], "Translation (draft)": t["translation"]}


def to_xlsx(rec: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    res = rec["result"]; ans = rec.get("answer") or {}
    wb = Workbook(); ws = wb.active; ws.title = "Answer"
    ws.append(["Question", rec["question"]]); ws.append(["Asked", rec.get("t", "")])
    ws.append(["Answer", ans.get("headline", "")]); ws.append(["", ans.get("answer", "")])
    for c in ans.get("claims") or []:
        ws.append(["Claim", c.get("text", ""), "evidence: " + str(c.get("evidence", ""))])
    for c in res["caveats"] + list(ans.get("caveats") or []):
        ws.append(["Caveat", c])
    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    def sheet(name, rows):
        s = wb.create_sheet(name[:31])
        if not rows:
            s.append(["(no rows)"]); return
        cols = list(rows[0]); s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True)
        for r in rows:
            s.append([", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v for v in (r.get(c) for c in cols)])
        s.freeze_panes = "A2"
    txdone = False
    for k, s in res["steps"].items():
        if s["type"] == "transactions" and not txdone:
            sheet(f"Evidence {k}", [_tx_flat(t) for t in s["rows"]]); txdone = True
        elif s["type"] == "table":
            sheet(f"{k} by {s['by']}", [{kk: v for kk, v in r.items()} for r in s["rows"]])
            if s.get("consistency"):
                sheet(f"{k} consistency", s["consistency"]); sheet(f"{k} by {s['over']}", s["matrix"])
        elif s["type"] == "network":
            sheet(f"{k} nodes", s["nodes"]); sheet(f"{k} edges", s["edges"])
        elif s["type"] == "query":
            sheet(f"{k} rows", [query.flat(r, s["unit"]) for r in s["rows"]])
            if s.get("groups"):
                sheet(f"{k} by {s['group_by']}", s["groups"])
    cs = wb.create_sheet("Coverage")
    for m in res["coverage"]["manuscripts"]:
        for kk, v in m.items():
            cs.append([kk, ", ".join(v) if isinstance(v, list) else v])
        cs.append([])
    ps = wb.create_sheet("Plan")
    import json
    ps.append([json.dumps(res["plan"], ensure_ascii=False, indent=1)])
    ps.column_dimensions["A"].width = 120
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def gephi_zip(net: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        s = io.StringIO(); w = csv.writer(s)
        w.writerow(["Id", "Label", "Type", "Role", "Rank", "Top", "Occupations", "Owed (n)", "Owed florins", "Owed lire",
                    "Owed to (n)", "Owed to florins", "Owed to lire", "Payments (n)", "Paid florins", "Paid lire", "Assumed"])
        for n in net["nodes"]:
            w.writerow([n["id"], n["label"], n["type"], n["role"], n["rank"], n["top"], "; ".join(n["occupations"]), n["owed_n"],
                        n["owed_florins"], n["owed_lire"], n["owed_to_n"], n["owed_to_florins"], n["owed_to_lire"], n["paid_n"], n["paid_florins"], n["paid_lire"], n["assumed"]])
        z.writestr("nodes.csv", "﻿" + s.getvalue())
        s = io.StringIO(); w = csv.writer(s)
        w.writerow(["Source", "Target", "Type", "Weight", "Florins owed", "Lire owed", "Payments", "Florins paid", "Lire paid", "Basis", "Evidence"])
        for e in net["edges"]:
            w.writerow([e["source"], e["target"], "Directed", max(1, e["n"]), e["florins"], e["lire"], e["paid_n"], e["paid_florins"], e["paid_lire"], e["basis"], " ".join(e["ids"])])
        z.writestr("edges.csv", "﻿" + s.getvalue())
        z.writestr("README.txt", "Directed: Source owed Target (debtor -> creditor). Gephi: File > Import spreadsheet > nodes.csv, then edges.csv.\n")
    return buf.getvalue()

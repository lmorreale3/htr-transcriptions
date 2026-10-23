"""Queries over one or more manuscripts → rows of lines, entries or tags; grouping; exports.

query = {
  "manuscripts": ["MPAP_4", …],          # default: all
  "statuses": [...] | null,              # default: each manuscript's setting
  "unit": "lines" | "entries" | "tags",
  "tag_types": [...],                    # unit "tags": which kinds of tag to list (default all)
  "approved_only": false,                # count only tags Laura approved
  "match": "all" | "any",
  "conditions": [ {field, op, value, scope?, attr?, which?}, … ]
}
Fields (ops):
  text (contains | not_contains | regex) — searches transcription, expanded text and translation; accents/abbreviation
       marks are ignored (an̄o matches "anno")
  person | place | occupation | institution | goods | any_tag   (exists | not_exists | is | contains)
       optional attr {"role": "debtor"} etc.; scope "line" (default) or "entry" (anywhere in the same entry)
  kind (is: debit | credit | balance | none)
  year (between [a, b]; which "modern" (default) | "written"),  month (is 1–12),  date (between ["1389-01-01", …], modern)
  page (between [first id, last id]),  book (contains),  folio (is)
  has_sum (is true|false),  florins (gte|lte),  lire (gte|lte, lire money in lire, decimals allowed)
  flagged (is true)
"""
from __future__ import annotations

import csv
import io
import re

from . import data, extract, source, tags as tagmod
from .textnorm import plain

TAGF = set(tagmod.TYPES) | {"any_tag"}


def _fold(s: str) -> str:
    return plain(s or "")


def _tag_match(tags, c, approved_only):
    typ = c["field"]; op = c.get("op", "exists"); v = _fold(str(c.get("value") or "")).strip()
    attr = c.get("attr") or {}
    ts = [t for t in tags if (typ == "any_tag" or t["type"] == typ) and (not approved_only or t["status"] == "approved")
          and all(str((t.get("attrs") or {}).get(k, "")).lower() == str(val).lower() for k, val in attr.items() if val)]
    if op == "exists":
        return bool(ts)
    if op == "not_exists":
        return not ts
    if op == "is":
        return any(_fold(t["canon"]) == v or _fold(t["name"]) == v for t in ts)
    if op == "contains":
        return any(v in _fold(t["canon"]) or v in _fold(t["name"]) or v in _fold(t.get("span", "")) for t in ts)
    raise ValueError(f"unknown op {op} for {typ}")


def _year(d, which):
    return d.get("year_written") if which == "written" else d.get("year_modern")


def _cond(row, c, ctx) -> bool:
    f = c.get("field"); op = c.get("op"); v = c.get("value")
    if f in TAGF:
        tags = ctx["entry_tags"].get(row.get("entry"), row["tags"]) if c.get("scope") == "entry" else row["tags"]
        return _tag_match(tags, c, ctx["approved_only"])
    if f == "text":
        hay = " ".join(_fold(row.get(k, "")) for k in ("text", "expanded", "translation", "sum", "date"))
        if op == "regex":
            return re.search(str(v), hay, re.I) is not None
        hit = _fold(str(v)) in hay
        return not hit if op == "not_contains" else hit
    if f == "kind":
        return (row.get("kind") or "none") == v
    if f in ("year", "date", "month"):
        ds = row.get("dates") or []
        if f == "month":
            return any(d.get("month") == int(v) for d in ds)
        a, b = (v or [None, None]) + [None] * (2 - len(v or []))
        for d in ds:
            x = _year(d, c.get("which")) if f == "year" else d.get("iso_modern")
            if x is None:
                continue
            if f == "year":
                if (a in (None, "") or x >= int(a)) and (b in (None, "") or x <= int(b)):
                    return True
            else:
                if (not a or x >= a) and (not b or x[:len(b)] <= b):
                    return True
        return False
    if f == "page":
        a, b = v
        o = ctx["page_order"].get((row["manuscript"], row["page"]))
        oa = ctx["page_order"].get((row["manuscript"], a), -1e9) if a else -1e9
        ob = ctx["page_order"].get((row["manuscript"], b), 1e9) if b else 1e9
        return o is not None and oa <= o <= ob
    if f == "book":
        return any(_fold(str(v)) in _fold(x.get("book_resolved") or "") for x in row.get("refs") or [])
    if f == "folio":
        return any(int(v) in x["folios"] for x in row.get("refs") or [])
    a = row.get("amount_final")
    if f == "has_sum":
        return bool(a) == (v in (True, "true", "yes", 1))
    if f in ("florins", "lire"):
        if not a:
            return False
        x = a.get("florins") if f == "florins" else a.get("lire_decimal")
        if x is None:
            return False
        return x >= float(v) if op == "gte" else x <= float(v)
    if f == "flagged":
        return bool(row.get("flag"))
    raise ValueError(f"unknown field {f}")


def _collect(q):
    mss = q.get("manuscripts") or [c["id"] for c in source.collections()]
    recs, order = [], {}
    for cid in mss:
        st = q.get("statuses") or data.settings(cid)["statuses"]
        for i, p in enumerate(source.pages(cid)):
            order[(cid, p["id"])] = i
        recs += extract.records(cid, st)
    return recs, order


def _entries(recs):
    groups = {}
    for r in recs:
        if r.get("entry"):
            groups.setdefault(r["entry"], []).append(r)
    return groups


def run(q: dict) -> dict:
    recs, order = _collect(q)
    groups = _entries(recs)
    ctx = {"approved_only": bool(q.get("approved_only")), "page_order": order,
           "entry_tags": {e: [t for r in rs for t in r["tags"]] for e, rs in groups.items()}}
    conds = q.get("conditions") or []
    anym = q.get("match") == "any"
    unit = q.get("unit", "lines")

    def ok(row):
        if not conds:
            return True
        res = (_cond(row, c, ctx) for c in conds)
        return any(res) if anym else all(res)

    if unit == "entries":
        rows = [entry_row(rs) for e, rs in groups.items()]
        rows = [r for r in rows if ok(r)]
    else:
        rows = [r for r in recs if ok(r)]
    if unit == "tags":
        types = set(q.get("tag_types") or tagmod.TYPES)
        # a condition on a kind of tag ("person is X") also picks which tags of that kind are listed
        own = [c for c in conds if c.get("field") in tagmod.TYPES and c.get("op") in ("is", "contains")]

        def keep(t):
            mine = [c for c in own if c["field"] == t["type"]]
            return not mine or any(_tag_match([t], c, ctx["approved_only"]) for c in mine)
        rows = [tag_row(r, t) for r in rows for t in r["tags"]
                if t["type"] in types and (not ctx["approved_only"] or t["status"] == "approved") and keep(t)]
    return {"unit": unit, "rows": rows, "count": len(rows), "lines_searched": len(recs)}


def entry_row(rs):
    first = rs[0]
    tags = [t for r in rs for t in r["tags"]]
    amts = [r["amount_final"] for r in rs if r.get("amount_final")]
    return {"key": first["entry"], "entry": first["entry"], "manuscript": first["manuscript"], "page": first["page"],
            "line": first["line"], "last_line": rs[-1]["line"], "line_ids": ",".join(r["line_ids"] for r in rs),
            "text": " / ".join(r["text"] for r in rs), "expanded": " / ".join(r["expanded"] for r in rs if r["expanded"]),
            "translation": " ".join(r["translation"] for r in rs if r["translation"]),
            "sum": " / ".join(r["sum"] for r in rs if r["sum"]), "date": "",
            "kind": next((r["kind"] for r in rs if r["kind"] in ("debit", "balance")), first["kind"]),
            "tags": tags, "dates": [d for r in rs for d in r["dates"]], "refs": [x for r in rs for x in r["refs"]],
            "amount_final": amts[0] if len(amts) == 1 else None, "amounts": amts, "flag": any(r["flag"] for r in rs),
            "status": first["status"], "lines": len(rs)}


def tag_row(r, t):
    return {"key": r["key"] + "|" + t["id"], "line_key": r["key"], "manuscript": r["manuscript"], "page": r["page"],
            "line": r["line"], "entry": r.get("entry"), "type": t["type"], "name": t["canon"], "name_as_tagged": t["name"],
            "span": t.get("span", ""), "attrs": t.get("attrs") or {}, "tag_status": t["status"], "text": r["text"],
            "translation": r.get("translation", ""), "dates": r.get("dates") or [], "kind": r.get("kind"),
            "line_ids": r["line_ids"]}


GROUPS = ["person", "place", "occupation", "institution", "goods", "role", "year_modern", "year_written", "month",
          "kind", "book", "page", "manuscript"]


def group(rows: list[dict], by: str, unit: str) -> list[dict]:
    """Counts per value: [{value, rows, first, last}] (first/last = earliest/latest modern date seen)."""
    buckets = {}
    for r in rows:
        if by in tagmod.TYPES:
            vals = {r["name"]} if unit == "tags" and r["type"] == by else (
                {t["canon"] for t in r.get("tags", []) if t["type"] == by} if unit != "tags" else set())
        elif by == "role":
            vals = {r["attrs"].get("role")} if unit == "tags" else {(t.get("attrs") or {}).get("role") for t in r.get("tags", []) if t["type"] == "person"}
            vals.discard(None)
        elif by in ("year_modern", "year_written", "month"):
            k = by if by != "month" else "month_name"
            vals = {d.get(k) for d in r.get("dates") or [] if d.get(k)}
        elif by == "book":
            vals = {x.get("book_resolved") for x in r.get("refs") or []}
        else:
            vals = {r.get(by) or "—"}
        for v in vals:
            b = buckets.setdefault(v, {"value": v, "rows": 0, "first": None, "last": None})
            b["rows"] += 1
            for d in r.get("dates") or []:
                iso = d.get("iso_modern")
                if iso:
                    b["first"] = min(b["first"] or iso, iso); b["last"] = max(b["last"] or iso, iso)
    return sorted(buckets.values(), key=lambda b: (-b["rows"], str(b["value"])))


# ------------------------------------------------------------------ flat rows for export
def _names(r, typ):
    return "; ".join(dict.fromkeys(t["canon"] for t in r.get("tags", []) if t["type"] == typ))


def flat(row: dict, unit: str) -> dict:
    ds = row.get("dates") or []
    common = {
        "Manuscript": row["manuscript"], "Page": row["page"], "Line": row["line"], "Entry": row.get("entry") or "",
        "Dates (modern)": "; ".join(d.get("iso_modern") or "" for d in ds),
        "Dates (as written, year)": "; ".join(f"{d.get('day') or ''} {d.get('month_name')} {d.get('year_written') or '?'}".strip() for d in ds),
        "Text (semi-diplomatic)": row["text"],
    }
    if unit == "tags":
        return {**common, "Type": row["type"], "Name": row["name"], "As written": row["span"],
                "Role": row["attrs"].get("role", ""), "Title": row["attrs"].get("title", ""),
                "Relation": row["attrs"].get("relation", ""), "Of": row["attrs"].get("of", ""),
                "Kind": row["attrs"].get("kind", ""), "Quantity": row["attrs"].get("quantity", ""),
                "Tag status": row["tag_status"], "Translation (draft)": row.get("translation", "")}
    a = row.get("amount_final")
    return {**common, "Expanded": row.get("expanded", ""), "Translation (draft)": row.get("translation", ""),
            "Entry kind": {"debit": "de dare", "credit": "de avere / anno dato", "balance": "resta"}.get(row.get("kind"), ""),
            "People": _names(row, "person"), "Places": _names(row, "place"), "Occupations": _names(row, "occupation"),
            "Institutions": _names(row, "institution"), "Goods": _names(row, "goods"),
            "References": "; ".join(f"{x.get('book_resolved') or ''} c. {'/'.join(map(str, x['folios']))}".strip() for x in row.get("refs") or []),
            "Sum (as written)": row.get("sum", ""), "Sum (read)": a["display"] if a else ""}


def to_csv(rows, unit) -> bytes:
    flat_rows = [flat(r, unit) for r in rows]
    buf = io.StringIO()
    if flat_rows:
        w = csv.DictWriter(buf, fieldnames=list(flat_rows[0])); w.writeheader(); w.writerows(flat_rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows, unit, about: list[tuple[str, str]], groups: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook(); ws = wb.active; ws.title = "Results"
    flat_rows = [flat(r, unit) for r in rows]
    cols = list(flat_rows[0]) if flat_rows else ["(no results)"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for fr in flat_rows:
        ws.append([fr[c] for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 60 if c.startswith(("Text", "Expanded", "Transl")) else 18
    ws.freeze_panes = "A2"
    for name, g in (groups or {}).items():
        gs = wb.create_sheet(f"By {name}"[:31]); gs.append([name, "rows", "first date", "last date"])
        for b in g:
            gs.append([str(b["value"]), b["rows"], b["first"] or "", b["last"] or ""])
    ab = wb.create_sheet("About")
    for k, v in about:
        ab.append([k, v])
    ab.column_dimensions["A"].width = 24; ab.column_dimensions["B"].width = 100
    for row in ab.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def describe(q: dict) -> str:
    parts = []
    for c in q.get("conditions") or []:
        v = c.get("value"); v = "–".join(map(str, v)) if isinstance(v, list) else v
        s = f"{c.get('field')} {c.get('op', '')} {v if v not in (None, '') else ''}".strip()
        if c.get("attr"):
            s += " (" + ", ".join(f"{k}={x}" for k, x in c["attr"].items() if x) + ")"
        if c.get("scope") == "entry":
            s += " [in the same entry]"
        parts.append(s)
    j = " OR " if q.get("match") == "any" else " AND "
    return (q.get("unit", "lines") + ": " + (j.join(parts) if parts else "everything")
            + " · manuscripts: " + ", ".join(q.get("manuscripts") or ["all"]))
